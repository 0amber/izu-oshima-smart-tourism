# First Stage収録カンペの docx / pdf を生成する
# 実行: /Users/shimadakoutaro/shoken/.venv/bin/python docs/submission/interview/build_kanpe.py
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).parent
DOCX = HERE / "2026-08-26-recording-kanpe.docx"
PDF = HERE / "2026-08-26-recording-kanpe.pdf"
HTML = HERE / "2026-08-26-recording-kanpe.html"

ACCENT = RGBColor(0x0E, 0x7A, 0x8A)  # 島テーマの青緑
DARK = RGBColor(0x1A, 0x1A, 0x1A)

SCRIPT_ROWS = [
    ("0:00–0:08", "1 表紙", "①",
     "その旅程、バスは本当にありますか？　私たちの作品は「大島スマートコース」です。▶"),
    ("0:08–0:22", "2 課題", "①",
     "生成AIに伊豆大島の旅程を頼むと、それらしい答えが返ってきます。でも、そのバスは存在しません。"
     "時刻表はPDF、GTFSはzip。AIは時刻表を読めないからです。▶"),
    ("0:22–0:38", "3 現実", "②",
     "島の現実はこうです。大きな荷物はバスで1個500円、混雑時は乗車拒否も。宿のチェックインは15時。"
     "三原山ラインは1日3往復で、午後に山へ上がる便はない。知らないと、初日が崩れます。▶"),
    ("0:38–0:49", "4 裏付け", "②",
     "来島者は年間約19万人。生成AI利用者の77.8%が旅行で活用する今、「正確な旅程を作るAI」はまだ空白です。▶"),
    ("0:49–1:04", "5 解決", "②",
     "そこで私たちは、公共交通オープンデータGTFSを観光に必要な停留所だけの表に変換し、"
     "現地の暗黙知10項目と一緒にAIへ渡しました。ルールはひとつ、「時刻表にない便を作らない」。▶"),
    ("1:04–1:18", "6 アプリ", "③",
     "公開中のアプリでは、日付・港・荷物・コースを選ぶだけ。実在する53便だけで時刻入りの旅程を組み、"
     "地図と天気、AIガイドの音声解説、日英切替まで、スマホですぐ使えます。▶"),
    ("1:18–1:28", "7 実演", "③",
     "実演では、荷物ありなら椿園経由、身軽なら山頂へ直行。コースを変えれば、港での乗り継ぎも自動で探します。▶"),
    ("1:28–1:38", "8 技術", "③",
     "時刻は実在便のみ、全便に検証マーク。公開する表は識別子も座標も含まない要約で、ライセンスにも配慮しています。▶"),
    ("1:38–1:49", "9 効果", "①",
     "到着直後の半日を、観光の時間に変える。GTFSのある島や地方へ、そのまま横展開できます。"
     "運用はOSS、ダイヤ改定ごとに1コマンドで更新します。▶"),
    ("1:49–2:00", "10 締め", "①",
     "オープンデータ×ドメイン知識×AI。「大島スマートコース」、今すぐ使えます。"
     "以上、SBC.別班連携チームの発表でした。ありがとうございました。"),
]

ROLES = [
    ("話者① つかみ・締め", "志田さん(登壇候補)", "スライド1-2, 9-10", "冒頭5秒の掴みとインパクトの言語化"),
    ("話者② 島の現実・解決", "杉山さん(大島で勤務経験あり)", "スライド3-5", "現地経験者が語る説得力"),
    ("話者③ アプリ・技術", "島田さん(開発・PM)", "スライド6-8", "デモと技術の確かさ"),
    ("画面共有・送り", "話者③が兼任", "oshima-slides.pptx(10枚)", "セリフ末尾の「▶」が送りの合図"),
]

CRITERIA = [
    ("データ活用", "5・10", "大島バスGTFS-JPが核。気象庁予報・OSM・都観光統計など計7種"),
    ("アイデア力", "2・7", "「幻覚ゼロの旅程」×「荷物で組み替わる」— 乗換案内と生成AIの間の空白を突く"),
    ("技術力", "8", "復元不可能な要約・全便検証マーク・AI失敗時フォールバック・1コマンド再生成"),
    ("ソーシャルインパクト", "4・9", "年間約19万人、国も「二次交通は喫緊の課題」と明記。半日を取り戻し島しょへ横展開"),
    ("サービスデザイン", "6", "スマホ・QRで今すぐ使える。日英切替・音声読み上げ・天気連動"),
]

CUTS = [
    "スライド9「運用はOSS、ダイヤ改定ごとに1コマンドで更新します。」をカット(約5秒)",
    "スライド3「知らないと、初日が崩れます。」をカット(約3秒)",
    "スライド8 後半を「ライセンスにも配慮しています。」に短縮(約4秒)",
]

CHECK_BEFORE = [
    "3人で通し練習2回(ストップウォッチ計測。1回目は素で、2回目は本番想定)",
    "oshima-slides.pptx をスライドショーで最後まで送れるか確認",
    "デモURL確認: izu-oshima-smart-tourism.kotaroshimada38.workers.dev",
    "Zoom表示名を「チーム名+氏名」に。マイク・カメラ・回線チェック",
]
CHECK_JUST = [
    "画面共有→スライドショー開始。発表者ツールは共有画面に映さない(カンペは紙か別デバイス)",
    "PC・スマホの通知オフ(おやすみモード)",
    "話者順(①→②→③→①)とスライド送り担当を最終確認",
]
CHECK_DURING = [
    "昨年は司会が「次は〇〇の発表です」とチーム名を紹介してから開始。長い名乗りは不要(台本どおり)",
    "噛んでも止まらない・謝らない。録り直し可否は当日の運営案内に従う",
    "押したら上記カット順で削る。最後の「ありがとうございました」は必ず言い切る",
    "YouTubeで一般公開される。顔出し・実名は同意済みメンバーのみ。写真クレジット・OSM帰属はそのまま",
]

QA = [
    ("データは何を?",
     "核は大島バスGTFS-JP(公共交通オープンデータセンター)。ODPT承認済みの実データで、生成した時刻表は公式時刻表と"
     "完全一致を確認。ほかに気象庁予報JSON、東京都の観光統計、OpenStreetMapなど計7種。"),
    ("AIの幻覚対策は?",
     "旅程の時刻はGTFSに実在する53便だけ。AI(Claude)は確定済みの旅程JSONを引用して解説するのみ。"
     "全便に「時刻表と一致」の検証マークと便IDを表示。"),
    ("ライセンスは大丈夫?",
     "公開物は識別子・座標を含まない復元不可能な要約に限定しODPTライセンスに配慮。写真はCC素材で帰属明記、"
     "地図は© OpenStreetMap contributors。"),
    ("他の地域でも使える?",
     "GTFSが整備されている島・地方なら変換は1コマンドで再生成でき、そのまま横展開できる。"),
    ("誰がどう使う?",
     "旅行者がスマホ・QRで直接使うほか、観光協会や宿の案内窓口が旅程提案に使う運用を、"
     "NPO・Code forのネットワークで現地と協働して進める。"),
]


def set_jp_font(run, size=None, bold=None, color=None):
    run.font.name = "Hiragino Kaku Gothic ProN"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Kaku Gothic ProN")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def heading(doc, text, size=14):
    p = doc.add_paragraph()
    p.space_before = Pt(6)
    r = p.add_run(text)
    set_jp_font(r, size=size, bold=True, color=ACCENT)
    return p


def para(doc, text, size=10.5, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_jp_font(r, size=size, bold=bold, color=DARK)
    return p


def bullets(doc, items, size=10.5):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(it)
        set_jp_font(r, size=size, color=DARK)


def table(doc, headers, rows, widths, body_size=10.5, bold_cols=()):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.paragraphs[0].text = ""
        r = cell.paragraphs[0].add_run(h)
        set_jp_font(r, size=body_size, bold=True, color=ACCENT)
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            cell = t.rows[i].cells[j]
            cell.paragraphs[0].text = ""
            r = cell.paragraphs[0].add_run(v)
            set_jp_font(r, size=body_size, bold=(j in bold_cols), color=DARK)
    for j, w in enumerate(widths):
        for row in t.rows:
            row.cells[j].width = Cm(w)
    return t


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(1.5)
    sec.left_margin = sec.right_margin = Cm(1.6)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("First Stage収録カンペ — 2分プレゼン(3人版)")
    set_jp_font(r, size=17, bold=True, color=ACCENT)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("SBC.別班連携チーム『大島スマートコース』｜収録 2026年8月26日・Zoom・日本語｜作成 2026-08-25")
    set_jp_font(r, size=9.5, color=DARK)

    para(doc,
         "公式の位置づけは「各チーム2分間のプレゼンテーション収録」(質疑ではない)。収録映像はYouTubeで一般公開。"
         "6人の審査委員が動画等で審査し、24作品がFinal Stage(10/17)へ進出。", size=10, bold=True)

    heading(doc, "1. 役割分担(推奨。入れ替え自由)")
    table(doc, ["役", "担当(案)", "受け持ち", "ねらい"], ROLES, [3.4, 4.0, 4.6, 5.8], body_size=9.5)

    heading(doc, "2. 本番台本(通し約120秒・670字)　※「▶」=次のスライドへ")
    table(doc, ["時間", "スライド", "話者", "セリフ"],
          SCRIPT_ROWS, [2.2, 2.4, 1.2, 12.0], body_size=11, bold_cols=(3,))

    heading(doc, "時間が押したときのカット(この順に)", size=11)
    bullets(doc, CUTS, size=10)

    doc.add_page_break()

    heading(doc, "3. 審査基準への刺さり方(勝ち筋)")
    para(doc, "審査基準は5つ(配点非公開)。台本は全基準を1回以上踏んでいる。", size=10)
    table(doc, ["審査基準", "刺すスライド", "ひとこと武器"], CRITERIA, [3.6, 2.4, 11.8], body_size=9.5)
    para(doc,
         "2025年度作品集(132作品)の傾向: 行政課題解決賞は交通系(風ぐるま乗換案内)。交通×課題解決は評価実績のある領域。"
         "本作の差別化は「実在保証(幻覚ゼロ)」と「荷物という独自の計画軸」。多数の企画止まり作品との違いとして、"
         "スライド6の「公開中」「今すぐ使えます」は必ず言う。", size=10)
    para(doc,
         "昨年のFirst Stageアーカイブ(4時間44分・約130チーム連続再生)から: 審査委員は2分プレゼンを大量に連続視聴する。"
         "冒頭5秒の問いかけで掴む・数字は少数を繰り返す・早口で詰め込まない、が効く。ファイナリストYORUGATA(1:25:46〜)の"
         "構成は「名乗り→統計で裏付けた課題→デモ→メリット・横展開→『以上、〇〇の発表でした』の定型締め」で、本台本も同型。2025年度審査委員: 宮坂学(東京都副知事)・井原正博(GovTech東京CTO)・閑歳孝子(くふうカンパニー)・"
         "高野克己(都デジタルサービス局長)・田村賢哉(Eukarya)・西村真里子(HEART CATCH)。行政×民間×技術の混成なので、"
         "専門用語(GTFS等)は台本どおり一言説明を添える。", size=10)

    heading(doc, "4. 収録前チェックリスト")
    para(doc, "前日(今日)", size=10.5, bold=True)
    bullets(doc, CHECK_BEFORE, size=10)
    para(doc, "開始直前", size=10.5, bold=True)
    bullets(doc, CHECK_JUST, size=10)
    para(doc, "収録中の心得・公開への注意", size=10.5, bold=True)
    bullets(doc, CHECK_DURING, size=10)

    heading(doc, "5. もし質問・確認をされたら(30秒短答集)")
    table(doc, ["想定質問", "短答"], [(q, a) for q, a in QA], [3.6, 14.2], body_size=9.5)

    para(doc,
         "出典: odhackathon.metro.tokyo.lg.jp /judging/・/recruitment/(審査基準)・/collection/(2025作品集)、"
         "docs/research/2026-08-21-market-research.md、提出フォーム最終版、昨年First Stageアーカイブ(youtube.com/watch?v=Ux1xQvWWf60)。", size=8.5)

    doc.save(DOCX)
    print("wrote", DOCX)


def build_html():
    def rows(items):
        return "\n".join(
            f"<tr><td class='t'>{t}</td><td class='s'>{s}</td><td class='sp'>{sp}</td>"
            f"<td class='line'>{line}</td></tr>"
            for t, s, sp, line in items)

    def trows(items):
        return "\n".join("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>" for r in items)

    def lis(items):
        return "\n".join(f"<li>{i}</li>" for i in items)

    html = f"""<meta charset="utf-8">
<style>
  @page {{ size: A4; margin: 0; }}
  * {{ box-sizing: border-box; margin: 0; padding: 0; }}
  body {{ font-family: "Hiragino Kaku Gothic ProN", "Hiragino Sans", sans-serif;
         color: #1a1a1a; font-size: 10.5pt; line-height: 1.55; }}
  .page {{ width: 210mm; min-height: 297mm; padding: 13mm 14mm; page-break-after: always; }}
  .page:last-child {{ page-break-after: auto; }}
  h1 {{ font-size: 16pt; color: #0e7a8a; text-align: center; }}
  .sub {{ text-align: center; font-size: 9pt; color: #555; margin: 2mm 0 4mm; }}
  .note {{ background: #eef7f8; border-left: 3px solid #0e7a8a; padding: 2.5mm 3.5mm;
           font-weight: 600; font-size: 9.5pt; margin-bottom: 4mm; }}
  h2 {{ font-size: 12pt; color: #0e7a8a; margin: 5mm 0 2mm; }}
  h3 {{ font-size: 10.5pt; margin: 3mm 0 1mm; }}
  table {{ border-collapse: collapse; width: 100%; }}
  th, td {{ border: 1px solid #b9c7ca; padding: 1.6mm 2.2mm; vertical-align: top; text-align: left; }}
  th {{ background: #e6f2f4; color: #0e5966; font-size: 9pt; }}
  .script td {{ font-size: 10pt; }}
  .script td.line {{ font-size: 11.5pt; font-weight: 600; }}
  .script td.t, .script td.sp, .script td.s {{ white-space: nowrap; font-weight: 700; color: #0e7a8a; }}
  .script td.s {{ color: #444; font-weight: 600; }}
  .small td {{ font-size: 9pt; }}
  ul {{ margin: 1mm 0 2mm 6mm; }}
  li {{ margin-bottom: 0.8mm; }}
  .src {{ font-size: 8pt; color: #777; margin-top: 4mm; }}
</style>
<div class="page">
  <h1>First Stage収録カンペ — 2分プレゼン(3人版)</h1>
  <div class="sub">SBC.別班連携チーム『大島スマートコース』｜収録 2026年8月26日・Zoom・日本語｜作成 2026-08-25</div>
  <div class="note">公式の位置づけは「各チーム2分間のプレゼンテーション収録」(質疑ではない)。収録映像はYouTubeで一般公開。
  6人の審査委員が動画等で審査し、24作品がFinal Stage(10/17)へ進出。</div>
  <h2>1. 役割分担(推奨。入れ替え自由)</h2>
  <table class="small"><tr><th>役</th><th>担当(案)</th><th>受け持ち</th><th>ねらい</th></tr>{trows(ROLES)}</table>
  <h2>2. 本番台本(通し約120秒・670字)　※「▶」=次のスライドへ</h2>
  <table class="script"><tr><th>時間</th><th>スライド</th><th>話者</th><th>セリフ</th></tr>{rows(SCRIPT_ROWS)}</table>
  <h3>時間が押したときのカット(この順に)</h3>
  <ul>{lis(CUTS)}</ul>
</div>
<div class="page">
  <h2>3. 審査基準への刺さり方(勝ち筋)</h2>
  <p>審査基準は5つ(配点非公開)。台本は全基準を1回以上踏んでいる。</p>
  <table class="small"><tr><th>審査基準</th><th>刺すスライド</th><th>ひとこと武器</th></tr>{trows(CRITERIA)}</table>
  <p style="margin-top:2mm">2025年度作品集(132作品)の傾向: 行政課題解決賞は交通系(風ぐるま乗換案内)。
  交通×課題解決は評価実績のある領域。本作の差別化は「実在保証(幻覚ゼロ)」と「荷物という独自の計画軸」。
  多数の企画止まり作品との違いとして、スライド6の<b>「公開中」「今すぐ使えます」は必ず言う</b>。</p>
  <p style="margin-top:2mm">昨年のFirst Stageアーカイブ(4時間44分・約130チーム連続再生)から: 審査委員は2分プレゼンを
  大量に連続視聴する。<b>冒頭5秒の問いかけで掴む・数字は少数を繰り返す・早口で詰め込まない</b>が効く。
  ファイナリストYORUGATA(1:25:46〜)の構成は「名乗り→統計で裏付けた課題→デモ→メリット・横展開→
  『以上、〇〇の発表でした』の定型締め」で、本台本も同型。2025年度審査委員: 宮坂学(東京都副知事)・井原正博(GovTech東京CTO)・
  閑歳孝子(くふうカンパニー)・高野克己(都デジタルサービス局長)・田村賢哉(Eukarya)・西村真里子(HEART CATCH)。
  行政×民間×技術の混成なので、専門用語(GTFS等)は台本どおり一言説明を添える。</p>
  <h2>4. 収録前チェックリスト</h2>
  <h3>前日(今日)</h3><ul>{lis(CHECK_BEFORE)}</ul>
  <h3>開始直前</h3><ul>{lis(CHECK_JUST)}</ul>
  <h3>収録中の心得・公開への注意</h3><ul>{lis(CHECK_DURING)}</ul>
  <h2>5. もし質問・確認をされたら(30秒短答集)</h2>
  <table class="small"><tr><th>想定質問</th><th>短答</th></tr>{trows(QA)}</table>
  <div class="src">出典: odhackathon.metro.tokyo.lg.jp /judging/・/recruitment/(審査基準)・/collection/(2025作品集)、
  docs/research/2026-08-21-market-research.md、提出フォーム最終版、昨年First Stageアーカイブ(youtube.com/watch?v=Ux1xQvWWf60)。</div>
</div>
"""
    HTML.write_text(html, encoding="utf-8")
    print("wrote", HTML)


def build_pdf():
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        page.goto(f"file://{HTML.absolute()}", wait_until="networkidle")
        page.wait_for_timeout(1500)
        page.pdf(path=str(PDF), width="210mm", height="297mm", print_background=True,
                 margin={"top": "0", "bottom": "0", "left": "0", "right": "0"})
        browser.close()
    print("wrote", PDF)


if __name__ == "__main__":
    build_docx()
    build_html()
    build_pdf()
