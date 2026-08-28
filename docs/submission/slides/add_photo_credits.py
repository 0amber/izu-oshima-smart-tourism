# oshima-slides.pptx に写真クレジットを追記する(スライド6・7=画面内写真、スライド10=一覧)
# 再実行可: 既存のクレジットボックスを消してから入れ直す
# 実行: /Users/shimadakoutaro/shoken/.venv/bin/python docs/submission/slides/add_photo_credits.py
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

HERE = Path(__file__).parent
PPTX = HERE / "oshima-slides.pptx"

GRAY = RGBColor(0x8A, 0x8A, 0x8A)      # 明背景(スライド6・7)用
LIGHT = RGBColor(0xB9, 0xCD, 0xD2)     # 暗背景(スライド10)用

# slide番号 -> (テキスト, 文字色, 幅比率)  幅を狭めてQR等と重ねない
CREDITS = {
    6: ("画面内写真: Guilhem Vellut / CC BY 2.0 (Wikimedia Commons)", GRAY, 0.94),
    7: ("画面内写真: Kentaro Ohno / CC BY 2.0・yano / CC BY-SA 3.0・Yoshi Canopus・(WT-shared) Shoestring・"
        "Tsuyoshi chiba / CC BY-SA 4.0 (Wikimedia Commons)／地図: © OpenStreetMap contributors", GRAY, 0.94),
    10: ("写真クレジット(Wikimedia Commons): Guilhem Vellut・Kentaro Ohno (CC BY 2.0)／lienyuan lee・alonfloc (CC BY 3.0)／"
         "Cassiopeia sweet・yano (CC BY-SA 3.0)／Yoshi Canopus・(WT-shared) Shoestring・Tsuyoshi chiba (CC BY-SA 4.0)／"
         "地図: © OpenStreetMap contributors", LIGHT, 0.72),
}

MARKS = ("画面内写真:", "写真クレジット(Wikimedia")


def remove_old(slide):
    for sh in list(slide.shapes):
        if sh.has_text_frame and any(m in sh.text_frame.text for m in MARKS):
            sh._element.getparent().remove(sh._element)


def main():
    prs = Presentation(PPTX)
    sw, sh_h = prs.slide_width, prs.slide_height
    for idx, slide in enumerate(prs.slides, start=1):
        if idx not in CREDITS:
            continue
        remove_old(slide)
        text, color, wratio = CREDITS[idx]
        box = slide.shapes.add_textbox(Emu(int(sw * 0.03)), Emu(int(sh_h * 0.875)),
                                       Emu(int(sw * wratio)), Emu(int(sh_h * 0.06)))
        tf = box.text_frame
        tf.word_wrap = True
        r = tf.paragraphs[0].add_run()
        r.text = text
        r.font.size = Pt(8)
        r.font.color.rgb = color
        r.font.name = "Hiragino Kaku Gothic ProN"
        print(f"slide {idx}: credit set")
    prs.save(PPTX)
    print("saved", PPTX)


if __name__ == "__main__":
    main()
